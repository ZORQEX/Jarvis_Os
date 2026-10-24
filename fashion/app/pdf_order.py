"""Factory order PDFs (proformas, order confirmations) → the same review and ledger path as an Excel file.

The item table is read with pdfplumber (text and cell positions, on this computer), each row's photo is cut out of
the page where it sits, and the result is written as a workbook in the layout the Excel reader already understands:
a title, the factory's own header and rows (photos in column B, the customer column unlabeled like in the factories'
Excel files), line totals as formulas when they agree with price × quantity, then the total, the factory's
percentage and the grand total as formulas. That workbook goes through ``factory_import.analyze_workbook``
unchanged: allocations («AA 40, Alina 38»), sizes, prices, sections, photos, the check against the factory's total.
A PDF without a readable table (a scan) is reported as such instead of guessed.
(Added for JARVIS, 2026-10-25: the brief asks for «Загрузка Excel и PDF … распознаёт товары, фотографии, размеры,
цены, клиентов и проценты».)
"""
from __future__ import annotations

import io
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

from app import factory_import as fi

ALLOCATION_HEADERS = {"client", "clients", "customer", "customers", "allocation", "клиент", "клиенты", "для кого",
                      "распределение", "cliente", "buyer", "for"}
PHOTO_HEADERS = fi.HEADER_TOKENS["image"]
TITLE_NOISE = re.compile(r"\b(proforma|pro-forma|order|confirmation|invoice|conferma|ordine|заказ|счёт|счет|page|стр)\b"
                         r"[\s:#№-]*[\w/-]*", re.I)
PERCENT = re.compile(r"(\d{1,2}(?:[.,]\d+)?)\s*%")


def number(text: str | None) -> Decimal | None:
    """«2 000,00», «2,000.00», «€ 1.500», «1500» → Decimal; anything else → None."""
    if re.search(r"[,;]\s", str(text or "")):
        return None                                               # «38, 40» is a list of sizes, not 38.40
    t = re.sub(r"[€$£\s  ]", "", str(text or ""))
    if not re.fullmatch(r"-?[\d.,']+", t) or not re.search(r"\d", t):
        return None
    t = t.replace("'", "")
    if "," in t and "." in t:
        t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
    elif "," in t:
        head, _, tail = t.rpartition(",")
        t = head.replace(",", "") + "." + tail if len(tail) in (1, 2) else t.replace(",", "")
    elif t.count(".") == 1 and len(t.rpartition(".")[2]) == 3 and len(t) > 4:
        t = t.replace(".", "")                                    # «1.500» = one thousand five hundred
    try:
        return Decimal(t)
    except InvalidOperation:
        return None


def _cell(text: str | None):
    t = re.sub(r"\s+", " ", str(text or "")).strip()
    if not t:
        return None
    n = number(t)
    if n is None:
        return t
    return int(n) if n == n.to_integral_value() else float(n)


def _title(page, top: float) -> str:
    above = page.crop((0, 0, page.width, max(1.0, top))).extract_text() or ""
    lines = [ln.strip() for ln in above.splitlines() if re.search(r"[A-Za-zА-Яа-я]{3}", ln)]
    return lines[0] if lines else ""


def _sheet_name(title: str, used: set[str], index: int) -> str:
    name = TITLE_NOISE.sub(" ", title)
    name = re.sub(r"[\[\]:*?/\\]", " ", name)
    name = re.sub(r"\s+", " ", name).strip(" -—–") or f"PDF {index}"
    name = name[:31].strip()
    base, n = name, 2
    while name in used:
        name = f"{base[:27]} {n}"
        n += 1
    used.add(name)
    return name


def _row_photo(page, images: list[dict], top: float, bottom: float, x0: float, x1: float):
    """The picture whose centre is inside this table row, cut out of the rendered page."""
    for im in images:
        cy, cx = (im["top"] + im["bottom"]) / 2, (im["x0"] + im["x1"]) / 2
        if top <= cy <= bottom and x0 - 2 <= cx <= x1 + 2:
            box = (max(0, im["x0"]), max(0, im["top"]), min(page.width, im["x1"]), min(page.height, im["bottom"]))
            try:
                picture = page.crop(box).to_image(resolution=150).original.convert("RGB")
            except Exception:   # noqa: BLE001 - a broken image never stops the rows
                return None
            out = io.BytesIO()
            picture.save(out, "PNG")
            return out.getvalue()
    return None


def _tables(page):
    found = page.find_tables()
    if not found:
        found = page.find_tables({"vertical_strategy": "text", "horizontal_strategy": "text"})
    return found


def read_pdf(path: Path) -> list[dict]:
    """Item tables of the PDF: [{title, header, rows: [{cells, photo}], tail: [lines]}]; a table continued on the next
    page (same header, or no header) joins the previous one."""
    import pdfplumber
    groups: list[dict] = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            images = list(page.images)
            for table in _tables(page):
                data = table.extract()
                if not data or len(data[0]) < 3:
                    continue
                score, _mapping, _runs = fi._header_score({i: c for i, c in enumerate(data[0]) if c})
                header = [re.sub(r"\s+", " ", c or "").strip() for c in data[0]]
                start = 1
                if score < 2:                                       # no header: a continuation of the last table
                    if not groups or len(groups[-1]["header"]) != len(header):
                        continue
                    group, start = groups[-1], 0
                elif groups and groups[-1]["header"] == header and table.bbox[1] < page.height * 0.25:
                    group = groups[-1]
                else:
                    group = {"title": _title(page, table.bbox[1]), "header": header, "rows": [], "tail": []}
                    groups.append(group)
                for i, cells in enumerate(data[start:], start):
                    if not any((c or "").strip() for c in cells):
                        continue
                    row = table.rows[i] if i < len(table.rows) else None
                    photo = _row_photo(page, images, row.bbox[1], row.bbox[3], table.bbox[0], table.bbox[2]) \
                        if row and row.bbox else None
                    group["rows"].append({"cells": cells, "photo": photo})
                below = page.crop((0, min(page.height, table.bbox[3]), page.width, page.height)).extract_text() or ""
                group["tail"] = [ln.strip() for ln in below.splitlines() if ln.strip()][:12]
    return groups


def pdf_to_workbook(pdf_path: Path, out_path: Path) -> dict:
    """Write the PDF's item tables as a factory-style workbook. → {"sheets", "rows", "photos"} (rows = 0: no table)."""
    groups = read_pdf(pdf_path)
    book = Workbook()
    book.remove(book.active)
    used, rows_total, photos = set(), 0, 0
    for n, g in enumerate(groups, 1):
        header = g["header"]
        keep = [i for i, h in enumerate(header) if fi.norm_label(h) not in PHOTO_HEADERS]
        sheet = book.create_sheet(_sheet_name(g["title"], used, n))
        sheet["B2"] = g["title"] or sheet.title
        col_of = {i: 3 + k for k, i in enumerate(keep)}             # PDF column → workbook column (C, D, …)
        for i in keep:
            label = header[i]
            sheet.cell(3, col_of[i], None if fi.norm_label(label) in ALLOCATION_HEADERS else label)
        _score, mapping, _ = fi._header_score({col_of[i]: header[i] for i in keep})
        pick = {k: v[0] for k, v in mapping.items()}
        first = row_no = 4
        for r in g["rows"]:
            cells = r["cells"]
            values = {col_of[i]: _cell(cells[i]) for i in keep if i < len(cells)}
            label_text = " ".join(str(v) for v in values.values() if isinstance(v, str))
            if fi.TOTAL_WORDS.search(label_text) and not (pick.get("article") and values.get(pick["article"])):
                continue                                             # the table's own total row: written below
            for col, v in values.items():
                sheet.cell(row_no, col, v)
            price, qty, total = (pick.get(k) for k in ("price", "qty", "total"))
            if price and qty and total:
                p, q, t = (values.get(c) for c in (price, qty, total))
                if all(isinstance(x, (int, float)) for x in (p, q, t)) and abs(p * q - t) < 0.01:
                    sheet.cell(row_no, total, f"={get_column_letter(price)}{row_no}*{get_column_letter(qty)}{row_no}")
            if r["photo"]:
                image = XLImage(io.BytesIO(r["photo"]))
                image.width = image.height = 60
                sheet.add_image(image, f"B{row_no}")
                photos += 1
            row_no += 1
            rows_total += 1
        last = row_no - 1
        total_col = pick.get("total")
        if total_col and last >= first:
            tl = get_column_letter(total_col)
            label_col = max(3, total_col - 1)
            sheet.cell(row_no, label_col, "Total")
            sheet.cell(row_no, total_col, f"=SUM({tl}{first}:{tl}{last})")
            rate = next((PERCENT.search(line) for line in g["tail"] if PERCENT.search(line)), None)
            if rate:
                fraction = number(rate.group(1)) / 100
                sheet.cell(row_no + 1, label_col, float(fraction))
                sheet.cell(row_no + 1, total_col, f"={tl}{row_no}*{get_column_letter(label_col)}{row_no + 1}")
                sheet.cell(row_no + 2, total_col, f"=SUM({tl}{row_no}:{tl}{row_no + 1})")
    if rows_total:
        book.save(out_path)
    return {"sheets": len(groups), "rows": rows_total, "photos": photos, "warnings": stated_totals(groups)}


def stated_totals(groups: list[dict]) -> list[str]:
    """The factory's own «Total» under the table vs the sum of its rows: a difference is said, never corrected."""
    out = []
    for g in groups:
        header = g["header"]
        _s, mapping, _ = fi._header_score({i: h for i, h in enumerate(header) if h})
        col = (mapping.get("total") or [None])[0]
        if col is None:
            continue
        rows = sum((number(r["cells"][col]) or 0) for r in g["rows"] if col < len(r["cells"])
                   and not fi.TOTAL_WORDS.search(" ".join(c or "" for c in r["cells"])))
        for line in g["tail"]:
            if fi.TOTAL_WORDS.search(line) and "%" not in line and not re.search(r"grand|итого с|всего с|с комисс", line, re.I):
                found = re.findall(r"[\d][\d\s.,']*\d", line)
                stated = number(found[-1]) if found else None
                if stated is not None and abs(stated - rows) > Decimal("0.01"):
                    out.append(f"{g['title'] or 'PDF'}: в PDF итог {stated}, а по строкам {rows} — проверьте строки.")
                break
    return out
