"""Excel and other tables as JARVIS's own tools (jarvis/tools/sheets.py): read, compute, create, change — and the
conversation using them instead of a Claude Code job."""
import datetime as dt
import shutil

import pytest
from openpyxl import Workbook, load_workbook

from conftest import ScriptedClaude, install, wait_done, wait_for
from jarvis.tasks import S
from jarvis.tools import sheets
from jarvis.tools.base import ToolContext, ToolInputError

SALES = [
    (dt.datetime(2026, 1, 15), "Иванов", "Ноутбук", 450000),
    (dt.datetime(2026, 3, 2), "Иванов", "Мышь", 8000),
    (dt.datetime(2026, 3, 20), "Петрова", "Ноутбук", 450000),
    (dt.datetime(2026, 3, 25), "Иванов", "Монитор", 120000),
    (dt.datetime(2026, 4, 1), "Сидоров", "Мышь", 16000),
]


def _ctx(core):
    return ToolContext(core.config, "t_test")


@pytest.fixture
def no_lo(monkeypatch):
    """Without LibreOffice: formula cells keep no value, everything else works the same."""
    monkeypatch.setattr(sheets, "soffice", lambda: None)
    sheets._cache.clear()


def _sales(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Продажи"
    ws.append(["Дата", "Менеджер", "Товар", "Сумма, ₸"])
    for r in SALES:
        ws.append(list(r))
    ws.append(["Итого", None, None, "=SUM(D2:D6)"])
    wb.create_sheet("План").append(["Месяц", "План"])
    wb.save(path)
    return path


def _budget(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Октябрь"
    rows = {1: ["Бюджет семьи — октябрь 2026"], 3: ["ДОХОДЫ"], 4: ["Зарплата", 650000], 5: ["Подработка", 80000],
            7: ["Итого доходы", "=SUM(B4:B5)"], 9: ["РАСХОДЫ"], 10: ["Аренда", 180000], 11: ["Продукты", 95000],
            12: ["Кафе", 18000], 13: ["Итого расходы", "=SUM(B10:B12)"], 15: ["Остаток", "=B7-B13"]}
    for r, vals in rows.items():
        for j, v in enumerate(vals, start=1):
            ws.cell(r, j).value = v
    wb.create_sheet("Итоги")["A1"] = "=Октябрь!B15*12"
    wb.save(path)
    return path


# ---------------------------------------------------------------- reading
async def test_read_finds_titles_data_rows_and_the_total(core, home, no_lo):
    p = _sales(home / "files" / "Продажи.xlsx")
    t = core.registry.tools["sheet_read"]
    assert t.spec.risk.value == "low"                        # reading never asks
    a = t.validate({"path": str(p)}, _ctx(core))
    r = await t.run(a, _ctx(core))
    text = r.data["text"]
    assert "заголовки в строке 1" in text and "D «Сумма, ₸»" in text
    assert "Строк с данными: 5 (строки 2–6)" in text and "Итоговая строка 7" in text
    assert "Иванов (3)" in text and r.data["sheets"] == ["Продажи", "План"]
    assert r.card == "table" and r.data["card"]["columns"][:2] == ["Дата", "Менеджер"] and r.untrusted
    assert (await t.verify(a, r, _ctx(core))).status == "verified"
    r = await t.run(t.validate({"path": str(p), "sheet": "2"}, _ctx(core)), _ctx(core))
    assert r.data["sheet"] == "План"


def test_header_row_tables_lists_and_forms():
    assert sheets.header_row([["Имя", "Телефон"], ["Анна", "+7 700"], ["Борис", "+7 701"]]) == 1     # all-text list
    assert sheets.header_row([["Отчёт"], [], ["Дата", "Сумма"], [dt.date(2026, 1, 1), 5]]) == 3
    form = [["Бюджет"], [], ["ДОХОДЫ"], ["Зарплата", 650000], ["Итого", 650000]]
    assert sheets.header_row(form) is None                  # a form in blocks: shown cell by cell


async def test_query_sums_groups_and_filters_exactly(core, home, no_lo):
    p = str(_sales(home / "files" / "Продажи.xlsx"))
    q = core.registry.tools["sheet_query"]

    async def run(**args):
        r = await q.run(q.validate({"path": p, **args}, _ctx(core)), _ctx(core))
        return r.data["rows"]
    assert await run(metrics=[{"fn": "sum", "column": "Сумма"}]) == [[1044000]]          # «Итого» is not data
    assert await run(group_by=["Товар"], metrics=[{"fn": "sum", "column": "D"}], sort={"by": 1}, limit=1) == \
        [["Ноутбук", 900000]]
    march = [{"column": "Менеджер", "op": "=", "value": "иванов"}, {"column": "Дата", "op": "=", "value": "2026-03"}]
    assert await run(where=march, metrics=[{"fn": "sum", "column": "Сумма"}, {"fn": "count"}]) == [[128000, 2]]
    assert await run(group_by=[{"column": "Дата", "by": "quarter"}], metrics=[{"fn": "count"}]) == \
        [["2026 Q1", 4], ["2026 Q2", 1]]
    assert await run(where=[{"column": "Сумма", "op": "between", "value": [10000, 200000]}],
                     metrics=[{"fn": "count_distinct", "column": "Менеджер"}]) == [[2]]
    with pytest.raises(ToolInputError, match="Столбцы: A «Дата»"):
        await run(metrics=[{"fn": "sum", "column": "Прибыль"}])


async def test_csv_with_semicolons_decimal_commas_and_cp1251(core, home):
    p = home / "files" / "Клиенты.csv"
    p.write_bytes("Имя;Город;Баланс\nАнна;Алматы;1 200,50\nБорис;Астана;300\nВера;алматы;99,5\n".encode("cp1251"))
    q = core.registry.tools["sheet_query"]
    a = q.validate({"path": str(p), "where": [{"column": "Город", "op": "=", "value": "Алматы"}],
                    "metrics": [{"fn": "count"}, {"fn": "sum", "column": "Баланс"}]}, _ctx(core))
    r = await q.run(a, _ctx(core))
    assert r.data["rows"] == [[2, 1300]] and r.data["data_rows"] == 3


async def test_read_file_shows_a_workbook_as_a_table(core, home, no_lo):
    p = _sales(home / "files" / "Продажи.xlsx")
    t = core.registry.tools["read_file"]
    r = await t.run(t.validate({"path": str(p)}, _ctx(core)), _ctx(core))
    assert "Строк с данными: 5" in r.data["text"]           # was «Формат .xlsx пока не поддерживается»


def test_reading_stays_inside_the_allowed_folders(core, home, tmp_path_factory):
    other = tmp_path_factory.mktemp("elsewhere")
    _sales(other / "x.xlsx")
    with pytest.raises(ToolInputError):
        core.registry.tools["sheet_read"].validate({"path": str(other / "x.xlsx")}, _ctx(core))
    (home / "files" / "notes.txt").write_text("x")
    with pytest.raises(ToolInputError, match="не таблица"):
        core.registry.tools["sheet_query"].validate({"path": str(home / "files" / "notes.txt")}, _ctx(core))


@pytest.mark.parametrize("text,decimal_comma,value", [
    ("31,037.55", False, 31037.55), ("1,234", False, 1234), ("1,234", True, 1.234), ("1.234,56", True, 1234.56),
    ("£1,200.00", False, 1200.0), ("(5 000)", True, -5000), ("1 200,50", True, 1200.5), ("007", True, "007"),
    ("04-Dec-25", False, dt.date(2025, 12, 4)), ("4 дек 2025 г.", True, dt.date(2025, 12, 4)),
    ("04/12/2025", False, dt.date(2025, 12, 4)), ("12/31/2025", False, dt.date(2025, 12, 31)),
    ("04.12.25", True, dt.date(2025, 12, 4)), ("Иванов", True, "Иванов"),
])
def test_numbers_and_dates_as_accounting_exports_write_them(text, decimal_comma, value):
    assert sheets._cell_from_text(text, decimal_comma) == value


async def test_an_english_ledger_export_sums_and_groups_by_month(core, home):
    """Seen in HM Treasury's «spending over £25,000» CSV: Windows-1252, «04-Dec-25», "31,037.55"."""
    p = home / "files" / "spend.csv"
    p.write_bytes('Entity,Date,Supplier, Amount \nDMO,04-Dec-25,Bank of England,"31,037.55"\n'
                  'HMT,16-Dec-25,Royal Mint – coins,"2,649,194.84"\nHMT,03-Jan-26,Royal Mint – coins,"£1,000.00"\n'
                  .encode("cp1252"))
    q = core.registry.tools["sheet_query"]
    r = await q.run(q.validate({"path": str(p), "group_by": [{"column": "Date", "by": "month"}],
                                "metrics": [{"fn": "sum", "column": "Amount"}]}, _ctx(core)), _ctx(core))
    assert r.data["rows"] == [["2025-12", 2680232.39], ["2026-01", 1000]]
    r = await q.run(q.validate({"path": str(p), "group_by": ["Supplier"], "metrics": [{"fn": "count"}]}, _ctx(core)), _ctx(core))
    assert ["Royal Mint – coins", 2] in r.data["rows"]                      # «–» read as cp1252, not as Cyrillic


def _form(path):
    """A Т-49-like unified form: a title, titles in two rows over merged boxes, the graph numbers, three employees,
    signatures."""
    wb = Workbook()
    ws = wb.active
    ws.title = "стр1"
    ws["AA1"] = "РАСЧЕТНО-ПЛАТЕЖНАЯ ВЕДОМОСТЬ"
    cols = {"A": ("Номер по порядку", None), "C": ("Табель- ный номер", None), "E": ("Начислено, руб.", "Оклад"),
            "G": (None, "всего"), "I": ("Удержано", "налог на доходы"), "K": ("Сумма, руб.", "к выплате"),
            "M": ("Деньги получил", "фамилия, инициалы")}
    for n, (c, (top, low)) in enumerate(cols.items(), start=1):
        nxt = sheets.col_letter(sheets.col_number(c) + 1)
        ws[f"{c}3"], ws[f"{c}4"], ws[f"{c}5"] = top, low, n
        for r in range(6, 9):
            ws.merge_cells(f"{c}{r}:{nxt}{r}")
    for r, (fio, oklad) in enumerate([("Иванова А.П.", 90000), ("Петров И.И.", 60000), ("Ким Д.", 45000)], start=6):
        ws[f"A{r}"], ws[f"C{r}"], ws[f"E{r}"], ws[f"G{r}"] = r - 5, f"00{r}", oklad, oklad
        ws[f"I{r}"], ws[f"K{r}"], ws[f"M{r}"] = round(oklad * 0.13), oklad - round(oklad * 0.13), fio
    ws["A11"], ws["K11"], ws["O11"], ws["S11"] = "Выплаты произвел", "кассир", "Семенова", "(подпись)"
    wb.save(path)
    return path


async def test_a_unified_form_is_a_table_under_its_graph_numbers(core, home, no_lo):
    p = _form(home / "files" / "Т-49.xlsx")
    t = sheets.table_of(sheets.load_book_sync(p).sheets[0])
    assert t.header == 5 and [r for r, _ in t.data] == [6, 7, 8]           # the signatures are not data
    assert t.names[6] == "Начислено, руб. / всего (гр. 4)" and t.names[2] == "Табельный номер (гр. 2)"
    q = core.registry.tools["sheet_query"]
    r = await q.run(q.validate({"path": str(p), "metrics": [{"fn": "sum", "column": "всего"},
                                                            {"fn": "sum", "column": "к выплате"}]}, _ctx(core)), _ctx(core))
    assert r.data["rows"] == [[195000, 169650]]
    r = await core.registry.tools["sheet_read"].run({"path": str(p), "from_row": 1, "max_rows": 60}, _ctx(core))
    assert "AA1=РАСЧЕТНО-ПЛАТЕЖНАЯ ВЕДОМОСТЬ" in r.data["text"] and "B «B»" not in r.data["text"]   # filled cells only
    card = r.data["card"]                     # the window's card: the titled graphs only, not «B C D E…»
    assert card["columns"][0] == "Номер по порядку (гр. 1)" and "B" not in card["columns"] and len(card["rows"]) == 3


async def test_a_form_is_filled_box_by_box(core, home, no_lo):
    p = _form(home / "files" / "Т-49.xlsx")
    e = core.registry.tools["sheet_edit"]
    a = e.validate({"path": str(p), "changes": [{"op": "set", "cell": "F6", "value": 95000},     # inside the E6:F6 box
                                                {"op": "insert_rows", "row": 9, "count": 1},
                                                {"op": "set", "cell": "M9", "value": "Новикова Е."}]}, _ctx(core))
    await e.run(a, _ctx(core))
    ws = load_workbook(p)["стр1"]
    assert ws["E6"].value == 95000 and ws["M9"].value == "Новикова Е."
    assert {"E9:F9", "M9:N9"} <= {str(m) for m in ws.merged_cells.ranges}  # the new row has the boxes of the row above
    with pytest.raises(ToolInputError, match="пишите в A9"):
        await e.run(e.validate({"path": str(p), "changes": [{"op": "insert_rows", "row": 9, "values": [[4, 5]]}]},
                               _ctx(core)), _ctx(core))


@pytest.mark.skipif(not sheets.soffice(), reason="LibreOffice is not installed")
async def test_an_old_format_blank_is_filled_into_a_new_xlsx(core, home):
    import subprocess
    sheets._cache.clear()
    src = _form(home / "Т-49.xlsx")
    subprocess.run([sheets.soffice(), f"-env:UserInstallation=file://{home}/lo", "--headless", "--convert-to", "ods",
                    "--outdir", str(home / "files"), str(src)], capture_output=True, timeout=120, check=True)
    ods = home / "files" / "Т-49.ods"
    before = ods.read_bytes()
    e = core.registry.tools["sheet_edit"]
    a = e.validate({"path": str(ods), "changes": [{"op": "set", "cell": "K6", "value": 80000}]}, _ctx(core))
    assert a["save_as"] == str(home / "files" / "Т-49.xlsx")
    r = await e.run(a, _ctx(core))
    assert ods.read_bytes() == before and not r.data["backup"]              # the blank itself is never changed
    assert load_workbook(a["save_as"])["стр1"]["K6"].value == 80000
    assert (await e.verify(a, r, _ctx(core))).status == "verified"


# ---------------------------------------------------------------- changing
def test_formulas_follow_inserted_and_deleted_rows_like_excel():
    f = sheets.shift_formula
    assert f("=SUM(B10:B14)", "Октябрь", 15, "Октябрь", 15, 1) == "=SUM(B10:B15)"   # an item added above the total
    assert f("=SUM(B10:B14)", "Октябрь", 5, "Октябрь", 15, 1) == "=SUM(B10:B14)"    # a formula above: no growth
    assert f("=B7-B15", "Октябрь", 17, "Октябрь", 15, 1) == "=B7-B16"
    assert f("=$B$14*2", "S", 1, "S", 10, 1) == "=$B$15*2"
    assert f("=Октябрь!B15*12", "Итоги", 1, "Октябрь", 15, 1) == "=Октябрь!B16*12"  # other sheets follow too
    assert f("=B15*12", "Итоги", 1, "Октябрь", 15, 1) == "=B15*12"                  # …but not their own cells
    assert f("=SUM(B10:B14)", "S", 20, "S", 12, -1) == "=SUM(B10:B13)"
    assert f("=B12+B13", "S", 20, "S", 12, -1) == "=#REF!+B12"


async def test_edit_inserts_a_row_moves_formulas_and_keeps_a_backup(core, home, no_lo):
    p = _budget(home / "files" / "Бюджет.xlsx")
    e = core.registry.tools["sheet_edit"]
    assert e.spec.risk.value == "moderate"                   # changing a file asks once
    a = e.validate({"path": str(p), "changes": [{"op": "insert_rows", "row": 13, "values": [["Связь", 6000]]}]},
                   _ctx(core))
    r = await e.run(a, _ctx(core))
    ws = load_workbook(p)["Октябрь"]
    assert ws["A13"].value == "Связь" and ws["B13"].value == 6000
    assert ws["B14"].value == "=SUM(B10:B13)" and ws["B16"].value == "=B7-B14"
    assert load_workbook(p)["Итоги"]["A1"].value == "=Октябрь!B16*12"
    assert r.data["backup"] and load_workbook(r.data["backup"])["Октябрь"]["A13"].value == "Итого расходы"
    assert (await e.verify(a, r, _ctx(core))).status == "verified"


def test_edit_never_loses_charts_in_place(core, home):
    from openpyxl.chart import BarChart, Reference
    p = _sales(home / "files" / "С диаграммой.xlsx")
    wb = load_workbook(p)
    ch = BarChart()
    ch.add_data(Reference(wb.active, min_col=4, min_row=1, max_row=6), titles_from_data=True)
    wb.active.add_chart(ch, "F2")
    wb.save(p)
    e = core.registry.tools["sheet_edit"]
    with pytest.raises(ToolInputError, match="диаграммы.*save_as"):
        e.validate({"path": str(p), "changes": [{"op": "set", "cell": "A1", "value": "x"}]}, _ctx(core))
    a = e.validate({"path": str(p), "save_as": "Копия", "changes": [{"op": "set", "cell": "A1", "value": "x"}]},
                   _ctx(core))
    assert a["save_as"] == str(home / "files" / "Копия.xlsx")


async def test_create_writes_into_jarvis_folder_and_refuses_to_overwrite(core, home, no_lo):
    c = core.registry.tools["sheet_create"]
    rows = [["Статья", "Сумма"], ["Зарплата", 650000], ["Аренда", -180000], ["Остаток", "=SUM(B2:B3)"]]
    a = c.validate({"path": "Бюджет на ноябрь", "sheets": [{"name": "Ноябрь", "rows": rows, "freeze": "A2"}]}, _ctx(core))
    assert a["path"] == str(home / "out" / "Бюджет на ноябрь.xlsx")
    r = await c.run(a, _ctx(core))
    ws = load_workbook(a["path"])["Ноябрь"]
    assert ws["B4"].value == "=SUM(B2:B3)" and ws["A1"].font.b and ws.freeze_panes == "A2"
    assert r.data["formulas"] == 1 and (await c.verify(a, r, _ctx(core))).status == "verified"
    with pytest.raises(ToolInputError, match="уже есть"):
        c.validate({"path": "Бюджет на ноябрь", "sheets": [{"rows": [[1]]}]}, _ctx(core))
    for bad in ("~/.ssh/keys.xlsx", "/etc/x.xlsx"):
        with pytest.raises(ToolInputError):
            c.validate({"path": bad, "sheets": [{"rows": [[1]]}]}, _ctx(core))


@pytest.mark.skipif(not sheets.soffice(), reason="LibreOffice is not installed")
async def test_libreoffice_recalculates_formulas_a_program_wrote(core, home):
    sheets._cache.clear()
    p = _budget(home / "files" / "Бюджет.xlsx")              # openpyxl stores formulas without their results
    assert sheets._missing_results(p)
    e = core.registry.tools["sheet_edit"]
    r = await e.run(e.validate({"path": str(p), "changes": [{"op": "insert_rows", "row": 13,
                                                             "values": [["Связь", 6000]]}]}, _ctx(core)), _ctx(core))
    assert r.data["values"]["Октябрь!B14 «Итого расходы»"] == 299000
    assert r.data["values"]["Октябрь!B16 «Остаток»"] == 431000 and not r.data["errors"]


@pytest.mark.skipif(not sheets.soffice(), reason="LibreOffice is not installed")
async def test_saved_files_keep_formula_results_for_other_programs(core, home):
    """openpyxl saves formulas without results: a phone preview or Telegram would show empty totals."""
    sheets._cache.clear()
    c = core.registry.tools["sheet_create"]
    rows = [["Сумма", 100, 250], ["Итого", "=SUM(B1:C1)", '="Итого: "&B2'], ["Дата", "=DATE(2026,11,1)", "=B1>C1"]]
    a = c.validate({"path": "Проверка", "sheets": [{"name": "Лист", "rows": rows}]}, _ctx(core))
    r = await c.run(a, _ctx(core))
    ws = load_workbook(a["path"], data_only=True)["Лист"]
    assert ws["B2"].value == 350 and ws["C2"].value == "Итого: 350" and ws["C3"].value in (False, 0)
    assert ws["B3"].value.date() == dt.date(2026, 11, 1)
    assert load_workbook(a["path"])["Лист"]["B2"].value == "=SUM(B1:C1)"       # the formula stays a formula
    assert (await c.verify(a, r, _ctx(core))).status == "verified"
    bad = c.validate({"path": "Ошибка", "sheets": [{"rows": [[1, "=A1/0"]]}]}, _ctx(core))
    r = await c.run(bad, _ctx(core))
    assert r.data["errors"] and "#DIV/0!" in r.data["errors"][0]
    assert (await c.verify(bad, r, _ctx(core))).status == "failed"


# ---------------------------------------------------------------- the conversation uses them
@pytest.mark.parametrize("phrase", [
    "Что в файле «Продажи 2026.xlsx» в папке Excel-тест?",                 # was a read-only Claude Code job
    "В файле «Бюджет семьи.xlsx» добавь в расходы строку «Связь» на 6 000",  # was a Claude Code change
    "Создай Excel-файл «Бюджет на ноябрь» с итогами формулами",
    "посчитай сумму по столбцу лог в файле отчет.csv",
])
def test_table_requests_go_to_the_conversation(phrase):
    from jarvis.agency.planner import heuristic_plan
    d, _ = heuristic_plan("t", phrase)
    assert [s.action for s in d.steps] == ["chat"]


def test_a_table_file_of_its_own_is_not_the_orders_ledger():
    from jarvis.director import _not_the_ledger
    assert _not_the_ledger("Сколько строк в файле «Клиенты.csv» и сколько клиентов из Алматы?")
    assert _not_the_ledger("В файле «Большая выгрузка.xlsx» сколько заказов отменено?")
    assert not _not_the_ledger("Загрузи файл фабрики «Заказ фабрики Dior SS27.xlsx» в заказы")
    assert not _not_the_ledger("сколько клиентов в заказе SS27")


@pytest.mark.parametrize("phrase,action", [
    ("Сравни платежи HMT за декабрь 2025 и март 2026: общая сумма в каждом месяце", "chat"),   # was a tool-less step
    ("сравни цены на iPhone 17 в Алматы и Астане", "chat"),
    ("Проанализируй этот документ", "analyze"),          # refers to the previous answer: it is given as data
])
async def test_a_comparison_needing_files_or_the_web_goes_to_the_conversation(core, phrase, action):
    d = await core.director.agency.plan(core.director, "t", phrase)
    assert [s.action for s in d.steps] == [action]


def test_a_one_off_step_never_shows_tool_call_text():
    from jarvis.providers.claude_chat import TOOL_MARKUP
    assert TOOL_MARKUP.search('Проверю.\n<invoke name="Bash">\n<parameter name="command">ls</parameter>')
    assert not TOOL_MARKUP.search("Сумма <b>выросла</b> на 5 % (данные < 2025)")


def test_code_about_tables_is_still_code():
    from jarvis.agency.planner import heuristic_plan
    d, _ = heuristic_plan("t", "исправь ошибку в скрипте, который читает excel в ~/Projects/jarvis")
    assert d.steps[0].action == "change_code"


async def test_a_question_about_a_workbook_needs_no_approval(core, home, no_lo):
    p = _sales(home / "files" / "Продажи.xlsx")
    g = ScriptedClaude(core.config, [[("sheet_query", {"path": str(p), "metrics": [{"fn": "sum", "column": "Сумма"}]}),
                                      "Общая выручка — 1 044 000 ₸."]])
    install(core, claude=g)
    tid = await core.director.submit("какая общая выручка в файле Продажи.xlsx?")
    t = await wait_done(core, tid)
    assert t["state"] == S.COMPLETED and not core.approvals.pending()
    assert [s["tool_id"] for s in t["steps"]] == ["agent_chat", "sheet_query"]
    out = g.tool_results[0]
    assert out["status"] == "ok" and out["data"]["rows"] == [[1044000]] and "1 044 000" in out["untrusted_content"]


async def test_changing_a_workbook_asks_first(core, home, no_lo):
    p = _budget(home / "files" / "Бюджет.xlsx")
    before = p.read_bytes()
    g = ScriptedClaude(core.config, [[("sheet_edit", {"path": str(p), "changes": [
        {"op": "insert_rows", "row": 13, "values": [["Связь", 6000]]}]}), "Добавил строку «Связь»."]])
    install(core, claude=g)
    tid = await core.director.submit("добавь в Бюджет.xlsx расход «Связь» 6000")
    await wait_for(lambda: core.approvals.pending())
    assert p.read_bytes() == before                         # nothing changes before the «да»
    core.director.decide(core.approvals.pending()[0]["id"], True)
    t = await wait_done(core, tid)
    assert t["state"] == S.COMPLETED and load_workbook(p)["Октябрь"]["A13"].value == "Связь"
    shutil.rmtree(home / ".jarvis-backups", ignore_errors=True)
